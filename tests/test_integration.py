"""
Integration tests for pss-export: SBML and TabularQual export.

Requires a running Neo4j instance with the PSS database.
Set connection details via environment variables or a .env file:

    MY_NEO4J_URI=bolt://localhost:7687
    MY_NEO4J_USER=neo4j
    MY_NEO4J_PASSWORD=password

Tests are automatically skipped if no database connection is available.
A small reaction subset is used (ABA pathway) to keep tests fast.
"""

import os
import pytest

from pss_export import GraphDB
from pss_export import PSSAdapter

TEST_PATHWAYS = ["Hormone - Abscisic acid (ABA)"]

# ---------------------------------------------------------------------------
# Fixtures
# ---------------------------------------------------------------------------

@pytest.fixture(scope="session")
def db():
    """Connect to Neo4j once per session; skip all tests if unavailable."""
    try:
        graph_db = GraphDB(
            uri=os.getenv("MY_NEO4J_URI", "bolt://localhost:7687"),
            user=os.getenv("MY_NEO4J_USER", "neo4j"),
            pwd=os.getenv("MY_NEO4J_PASSWORD", "password"),
        )
        yield graph_db
        graph_db.close()
    except Exception as e:
        pytest.skip(f"Neo4j unavailable: {e}")


@pytest.fixture(scope="session")
def adapter(db):
    """PSSAdapter with a small ABA pathway subset collected once per session."""
    adapter = PSSAdapter(
        db,
        model_id="test_model",
        model_name="Test Export Model",
    )
    adapter.collect_reactions(
        access="public",
        pathways=TEST_PATHWAYS,
    )
    return adapter


# ---------------------------------------------------------------------------
# Collection
# ---------------------------------------------------------------------------

class TestCollectReactions:

    def test_reactions_collected(self, adapter):
        assert len(adapter.reaction_ids) > 0

    def test_reactions_dict_matches_ids(self, adapter):
        assert set(adapter.reaction_ids) == set(adapter.reactions.keys())

    def test_each_reaction_has_type(self, adapter):
        for rxn in adapter.reactions.values():
            assert rxn.reaction_type is not None

    def test_export_datetime_set(self, adapter):
        assert adapter.export_datetime is not None


# ---------------------------------------------------------------------------
# SBML export
# ---------------------------------------------------------------------------

class TestSBMLExport:

    @pytest.fixture(scope="class")
    def sbml_string(self, adapter):
        """Export to string (no file written)."""
        return adapter.create_sbml(filename=None)

    @pytest.fixture(scope="class")
    def sbml_file(self, adapter, tmp_path_factory):
        """Export to a temp file."""
        out = tmp_path_factory.mktemp("sbml") / "test_export.sbml"
        adapter.create_sbml(filename=str(out))
        return out

    def test_sbml_string_not_empty(self, sbml_string):
        assert sbml_string is not None
        assert len(sbml_string) > 0

    def test_sbml_string_is_valid_xml(self, sbml_string):
        import xml.etree.ElementTree as ET
        ET.fromstring(sbml_string)  # raises if invalid

    def test_sbml_string_declares_sbml(self, sbml_string):
        assert "sbml" in sbml_string.lower()

    def test_sbml_file_created(self, sbml_file):
        assert sbml_file.exists()
        assert sbml_file.stat().st_size > 0

    def test_sbml_file_content_matches_string(self, sbml_string, sbml_file):
        assert sbml_file.read_text() == sbml_string

    def test_sbml_contains_listof_reactions(self, sbml_string):
        assert "listOfReactions" in sbml_string

    def test_sbml_contains_listof_species(self, sbml_string):
        assert "listOfSpecies" in sbml_string

    def test_sbml_contains_compartments(self, sbml_string):
        assert "listOfCompartments" in sbml_string


# ---------------------------------------------------------------------------
# TabularQual export
# ---------------------------------------------------------------------------
class TestTabularQualExport:

    @pytest.fixture(scope="class")
    def tabularqual_file(self, adapter, tmp_path_factory):
        """Export to a temp .xlsx file.

        Note: write_spreadsheet (tabularqual) silently appends .xlsx to any
        path that doesn't already have an Excel extension, so passing .tsv
        would produce test_export.tsv.xlsx — a different path than expected.
        Use .xlsx directly for a predictable output path.
        """
        out = tmp_path_factory.mktemp("tabularqual") / "test_export.xlsx"
        adapter.create_tabulrqual(filename=str(out))
        return out

    def test_file_created(self, tabularqual_file):
        assert tabularqual_file.exists()
        assert tabularqual_file.stat().st_size > 0

    def test_file_is_xlsx(self, tabularqual_file):
        # xlsx files are ZIP archives; check magic bytes PK\x03\x04
        with open(tabularqual_file, "rb") as f:
            assert f.read(4) == b"PK\x03\x04"

    def test_file_has_expected_sheets(self, tabularqual_file):
        import openpyxl
        wb = openpyxl.load_workbook(tabularqual_file)
        assert "Species" in wb.sheetnames
        assert "Transitions" in wb.sheetnames



# ---------------------------------------------------------------------------
# Pathway filter
# ---------------------------------------------------------------------------

class TestPathwayFilter:

    def test_reaction_subset_smaller_than_full(self, db):
        """Filtering by pathway should return fewer reactions than no filter."""
        full = PSSAdapter(db)
        full.collect_reactions(access="public")

        subset = PSSAdapter(db)
        subset.collect_reactions(access="public", pathways=TEST_PATHWAYS)

        assert len(subset.reaction_ids) < len(full.reaction_ids)
